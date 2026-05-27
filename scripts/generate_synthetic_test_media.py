"""Generate privacy-safe synthetic lecture fixtures for Phase 10 demos.

The committed fixture source lives under fixtures/synthetic_media/source. This
script expands it into generated media, slide, caption, transcript, and manifest
files that can be uploaded through the app without using private recordings.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import html
import json
import math
import shutil
import struct
import subprocess
import wave
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_ROOT = REPO_ROOT / "fixtures" / "synthetic_media"
SOURCE_DIR = FIXTURE_ROOT / "source"
DEFAULT_OUTPUT_DIR = FIXTURE_ROOT / "generated"

MANIFEST_SOURCE = SOURCE_DIR / "synthetic_lecture_manifest.json"
TRANSCRIPT_SOURCE = SOURCE_DIR / "synthetic_lecture_transcript.json"
SLIDES_SOURCE = SOURCE_DIR / "synthetic_lecture_slides.json"


def load_source_fixtures() -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    return (
        _read_json(MANIFEST_SOURCE),
        _read_json(TRANSCRIPT_SOURCE),
        _read_json(SLIDES_SOURCE),
    )


def build_word_level_transcript(transcript: dict[str, Any]) -> dict[str, Any]:
    expanded = copy.deepcopy(transcript)
    words: list[dict[str, Any]] = []

    for segment in expanded.get("segments", []):
        segment_words = _tokenize_words(segment.get("text", ""))
        if not segment_words:
            continue

        start = float(segment["start_time"])
        end = float(segment["end_time"])
        duration = max(0.1, end - start)
        step = duration / len(segment_words)
        word_items = []

        for index, word in enumerate(segment_words):
            word_start = round(start + index * step + min(0.05, step * 0.1), 3)
            word_end = round(min(end, start + (index + 1) * step - min(0.05, step * 0.1)), 3)
            item = {
                "word": word,
                "start": word_start,
                "end": max(word_start, word_end),
                "confidence": 1.0,
                "speaker": segment.get("speaker", "lecturer"),
                "segment_id": segment["id"],
            }
            word_items.append(item)
            words.append(item)

        segment["words"] = word_items

    expanded["words"] = words
    expanded["metadata"] = {
        "generated_by": "scripts/generate_synthetic_test_media.py",
        "word_timing_strategy": "evenly_distributed_within_segments",
        "contains_private_recordings": False,
    }
    return expanded


def write_metadata_fixtures(
    output_dir: Path,
    *,
    force: bool = False,
    skip_pptx: bool = False,
) -> dict[str, Any]:
    manifest, transcript_source, slides = load_source_fixtures()
    transcript = build_word_level_transcript(transcript_source)

    transcript_dir = output_dir / "transcripts"
    slides_dir = output_dir / "slides"
    transcript_path = transcript_dir / "synthetic_lecture_transcript.json"
    srt_path = transcript_dir / "synthetic_lecture.srt"
    vtt_path = transcript_dir / "synthetic_lecture.vtt"
    slide_svg_paths = write_slide_svgs(slides, slides_dir, force=force)
    pptx_path = None if skip_pptx else write_pptx_deck(slides, slides_dir / "synthetic_lecture_slides.pptx", force=force)

    _write_json(transcript_path, transcript, force=force)
    srt_content = build_srt(transcript["segments"])
    _write_text(srt_path, srt_content, force=force)
    _write_text(vtt_path, "WEBVTT\n\n" + srt_content.replace(",", "."), force=force)

    generated_manifest = copy.deepcopy(manifest)
    generated_manifest["generated_at"] = datetime.now(timezone.utc).isoformat()
    generated_manifest["source_fixture_paths"] = {
        "manifest": _relative_to_repo(MANIFEST_SOURCE),
        "transcript": _relative_to_repo(TRANSCRIPT_SOURCE),
        "slides": _relative_to_repo(SLIDES_SOURCE),
    }
    generated_manifest["generated_sidecars"] = {
        "srt": _relative_to_output(srt_path, output_dir),
        "vtt": _relative_to_output(vtt_path, output_dir),
        "slide_svgs": [_relative_to_output(path, output_dir) for path in slide_svg_paths],
        "pptx": _relative_to_output(pptx_path, output_dir) if pptx_path else None,
    }

    manifest_path = output_dir / "synthetic_lecture_manifest.json"
    _write_json(manifest_path, generated_manifest, force=force)
    return {
        "manifest_path": manifest_path,
        "transcript_path": transcript_path,
        "srt_path": srt_path,
        "vtt_path": vtt_path,
        "slide_svg_paths": slide_svg_paths,
        "pptx_path": pptx_path,
    }


def write_media_fixtures(output_dir: Path, *, force: bool = False, ffmpeg_bin: str = "ffmpeg") -> dict[str, Path]:
    manifest, _, _ = load_source_fixtures()
    duration = float(manifest["duration_seconds"])
    media_dir = output_dir / "media"
    media_dir.mkdir(parents=True, exist_ok=True)

    audio_path = media_dir / "separate_audio.wav"
    screen_path = media_dir / "screen_recording.mp4"
    camera_path = media_dir / "webcam_recording.mp4"
    lecture_path = media_dir / "lecture_video.mp4"

    if force or not audio_path.exists():
        write_synthetic_audio(audio_path, duration)

    resolved_ffmpeg = resolve_ffmpeg_binary(ffmpeg_bin)
    if not resolved_ffmpeg:
        raise RuntimeError(
            f"FFmpeg binary '{ffmpeg_bin}' was not found. Re-run with --metadata-only "
            "or install FFmpeg to generate MP4 media."
        )

    if force or not screen_path.exists():
        _run_ffmpeg(
            [
                resolved_ffmpeg,
                "-y",
                "-f",
                "lavfi",
                "-i",
                f"testsrc2=size=1280x720:rate=30:duration={duration}",
                "-vf",
                "drawbox=x=48:y=48:w=1184:h=624:color=white@0.08:t=fill,drawgrid=width=160:height=90:thickness=1:color=white@0.16",
                "-c:v",
                "libx264",
                "-preset",
                "veryfast",
                "-crf",
                "23",
                "-pix_fmt",
                "yuv420p",
                "-an",
                str(screen_path),
            ]
        )

    if force or not camera_path.exists():
        _run_ffmpeg(
            [
                resolved_ffmpeg,
                "-y",
                "-f",
                "lavfi",
                "-i",
                f"color=c=0x18202b:size=640x360:rate=30:duration={duration}",
                "-vf",
                "drawbox=x=250:y=54:w=140:h=140:color=0xf59e0b@0.85:t=fill,drawbox=x=205:y=210:w=230:h=112:color=0x38bdf8@0.7:t=fill,drawgrid=width=80:height=60:thickness=1:color=white@0.12",
                "-c:v",
                "libx264",
                "-preset",
                "veryfast",
                "-crf",
                "24",
                "-pix_fmt",
                "yuv420p",
                "-an",
                str(camera_path),
            ]
        )

    if force or not lecture_path.exists():
        _run_ffmpeg(
            [
                resolved_ffmpeg,
                "-y",
                "-i",
                str(screen_path),
                "-i",
                str(camera_path),
                "-i",
                str(audio_path),
                "-filter_complex",
                "[1:v]scale=320:180[cam];[0:v][cam]overlay=W-w-32:H-h-32:format=auto[v]",
                "-map",
                "[v]",
                "-map",
                "2:a",
                "-shortest",
                "-c:v",
                "libx264",
                "-preset",
                "veryfast",
                "-crf",
                "23",
                "-pix_fmt",
                "yuv420p",
                "-c:a",
                "aac",
                "-b:a",
                "128k",
                str(lecture_path),
            ]
        )

    return {
        "audio": audio_path,
        "screen": screen_path,
        "camera": camera_path,
        "lecture": lecture_path,
    }


def write_synthetic_audio(path: Path, duration_seconds: float, sample_rate: int = 44100) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    total_samples = int(duration_seconds * sample_rate)
    silence_windows = [(16.0, 16.35), (44.4, 48.2)]
    cue_beeps = [0.0, 12.0, 26.0, 42.0, 60.0]

    with wave.open(str(path), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(sample_rate)
        frames = bytearray()

        for index in range(total_samples):
            t = index / sample_rate
            if any(start <= t <= end for start, end in silence_windows):
                sample = 0
            else:
                base = 0.18 * math.sin(2 * math.pi * 220 * t)
                harmonic = 0.05 * math.sin(2 * math.pi * 440 * t)
                cue = 0.0
                if any(cue_start <= t <= cue_start + 0.16 for cue_start in cue_beeps):
                    cue = 0.35 * math.sin(2 * math.pi * 880 * t)
                sample = int(max(-1.0, min(1.0, base + harmonic + cue)) * 32767)
            frames.extend(struct.pack("<h", sample))

        audio.writeframes(frames)


def resolve_ffmpeg_binary(ffmpeg_bin: str = "ffmpeg") -> str | None:
    """Find FFmpeg on PATH or in common Winget install locations."""
    if shutil.which(ffmpeg_bin):
        return ffmpeg_bin

    candidate = Path(ffmpeg_bin)
    if candidate.exists():
        return str(candidate)

    local_app_data = Path.home() / "AppData" / "Local"
    winget_packages = local_app_data / "Microsoft" / "WinGet" / "Packages"
    if winget_packages.exists():
        matches = sorted(winget_packages.glob("Gyan.FFmpeg_*/*/bin/ffmpeg.exe"))
        if matches:
            return str(matches[-1])

    return None


def write_slide_svgs(slides: dict[str, Any], slides_dir: Path, *, force: bool = False) -> list[Path]:
    outputs = []
    for slide in slides.get("slides", []):
        path = slides_dir / f"slide_{int(slide['number']):02d}.svg"
        bullet_text = "".join(
            f'<text x="110" y="{330 + index * 58}" class="bullet">- {html.escape(bullet)}</text>'
            for index, bullet in enumerate(slide.get("bullets", []))
        )
        content = f"""<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">
  <rect width="1280" height="720" fill="#111827"/>
  <rect x="64" y="64" width="1152" height="592" rx="24" fill="#f8fafc"/>
  <rect x="64" y="64" width="1152" height="104" rx="24" fill="#0f766e"/>
  <text x="104" y="132" font-family="Arial, sans-serif" font-size="42" font-weight="700" fill="#ffffff">{html.escape(slide["title"])}</text>
  <text x="104" y="240" font-family="Arial, sans-serif" font-size="30" fill="#0f172a">{html.escape(slide["subtitle"])}</text>
  <g font-family="Arial, sans-serif" font-size="26" fill="#334155">{bullet_text}</g>
  <text x="1100" y="620" font-family="Arial, sans-serif" font-size="24" fill="#64748b">Slide {int(slide["number"])}</text>
</svg>
"""
        _write_text(path, content, force=force)
        outputs.append(path)
    return outputs


def write_pptx_deck(slides: dict[str, Any], path: Path, *, force: bool = False) -> Path | None:
    if path.exists() and not force:
        return path

    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return None

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    for slide_data in slides.get("slides", []):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        title = slide.shapes.title
        title.text = slide_data["title"]
        title.text_frame.paragraphs[0].font.size = Pt(36)

        subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.6), Inches(0.6))
        subtitle.text_frame.text = slide_data["subtitle"]
        subtitle.text_frame.paragraphs[0].font.size = Pt(22)

        body = slide.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(11.2), Inches(3.7))
        frame = body.text_frame
        frame.clear()
        for index, bullet in enumerate(slide_data.get("bullets", [])):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(20)

    presentation.save(path)
    return path


def build_srt(segments: list[dict[str, Any]]) -> str:
    blocks = []
    for index, segment in enumerate(segments, start=1):
        blocks.append(
            "\n".join(
                [
                    str(index),
                    f"{_srt_time(float(segment['start_time']))} --> {_srt_time(float(segment['end_time']))}",
                    segment["text"],
                ]
            )
        )
    return "\n\n".join(blocks) + "\n"


def describe_generation_plan(output_dir: Path) -> list[str]:
    return [
        str(output_dir / "synthetic_lecture_manifest.json"),
        str(output_dir / "media" / "lecture_video.mp4"),
        str(output_dir / "media" / "screen_recording.mp4"),
        str(output_dir / "media" / "webcam_recording.mp4"),
        str(output_dir / "media" / "separate_audio.wav"),
        str(output_dir / "slides" / "synthetic_lecture_slides.pptx"),
        str(output_dir / "transcripts" / "synthetic_lecture_transcript.json"),
        str(output_dir / "transcripts" / "synthetic_lecture.srt"),
        str(output_dir / "transcripts" / "synthetic_lecture.vtt"),
    ]


def checksum_manifest(paths: dict[str, Path]) -> dict[str, str]:
    checksums = {}
    for key, path in paths.items():
        if path.exists() and path.is_file():
            checksums[key] = hashlib.sha256(path.read_bytes()).hexdigest()
    return checksums


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate synthetic AI video editor test media.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--metadata-only", action="store_true", help="Skip FFmpeg media generation.")
    parser.add_argument("--force", action="store_true", help="Overwrite existing generated files.")
    parser.add_argument("--dry-run", action="store_true", help="Print the planned output files without writing.")
    parser.add_argument("--skip-pptx", action="store_true", help="Skip optional PPTX generation.")
    parser.add_argument("--ffmpeg-bin", default="ffmpeg", help="FFmpeg executable to use for MP4 generation.")
    args = parser.parse_args()

    if args.dry_run:
        print("Synthetic fixture generation plan:")
        for item in describe_generation_plan(args.output_dir):
            print(f"- {item}")
        return 0

    outputs = write_metadata_fixtures(args.output_dir, force=args.force, skip_pptx=args.skip_pptx)
    media_outputs = {}
    if not args.metadata_only:
        media_outputs = write_media_fixtures(args.output_dir, force=args.force, ffmpeg_bin=args.ffmpeg_bin)

    print(f"Synthetic fixtures generated in {args.output_dir}")
    print(f"Manifest: {outputs['manifest_path']}")
    for key, path in {**media_outputs, "transcript": outputs["transcript_path"]}.items():
        if path:
            print(f"{key}: {path}")
    return 0


def _read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: dict[str, Any], *, force: bool) -> None:
    _write_text(path, json.dumps(payload, indent=2, sort_keys=False) + "\n", force=force)


def _write_text(path: Path, content: str, *, force: bool) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and not force:
        return
    path.write_text(content, encoding="utf-8")


def _run_ffmpeg(args: list[str]) -> None:
    args = [args[0], "-hide_banner", "-loglevel", "error", *args[1:]]
    subprocess.run(args, check=True)


def _tokenize_words(text: str) -> list[str]:
    return [token.strip() for token in text.replace(",", "").replace(".", "").split() if token.strip()]


def _srt_time(seconds: float) -> str:
    millis = int(round(seconds * 1000))
    hours = millis // 3_600_000
    millis %= 3_600_000
    minutes = millis // 60_000
    millis %= 60_000
    secs = millis // 1000
    millis %= 1000
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def _relative_to_repo(path: Path) -> str:
    return path.resolve().relative_to(REPO_ROOT.resolve()).as_posix()


def _relative_to_output(path: Path | None, output_dir: Path) -> str | None:
    if path is None:
        return None
    return path.resolve().relative_to(output_dir.resolve()).as_posix()


if __name__ == "__main__":
    raise SystemExit(main())
